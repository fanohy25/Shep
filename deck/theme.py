"""
Système de design du deck Shep.
Palette, typographie, grille et helpers de rendu bas-niveau (python-pptx).
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# --------------------------------------------------------------------------
# PALETTE
# --------------------------------------------------------------------------
INK        = RGBColor(0x08, 0x14, 0x24)   # fond sombre principal
INK_SOFT   = RGBColor(0x10, 0x21, 0x36)   # carte sur fond sombre
NAVY       = RGBColor(0x0F, 0x2A, 0x47)
ACCENT     = RGBColor(0x2E, 0x7D, 0xFF)   # bleu signature
ACCENT_DK  = RGBColor(0x17, 0x54, 0xC4)
TEAL       = RGBColor(0x00, 0xC2, 0xA8)   # accent secondaire
AMBER      = RGBColor(0xF5, 0xA6, 0x23)
CORAL      = RGBColor(0xF0, 0x62, 0x5B)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
PAPER      = RGBColor(0xF7, 0xF9, 0xFC)   # fond clair
LINE_LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
SLATE      = RGBColor(0x5B, 0x6C, 0x81)   # texte secondaire clair
SLATE_DK   = RGBColor(0x93, 0xA6, 0xBD)   # texte secondaire sombre
GRAPHITE   = RGBColor(0x14, 0x21, 0x32)   # texte principal sur clair

# --------------------------------------------------------------------------
# TYPOGRAPHIE
# --------------------------------------------------------------------------
FONT       = "Segoe UI"
FONT_LIGHT = "Segoe UI Light"
FONT_SEMI  = "Segoe UI Semibold"

# --------------------------------------------------------------------------
# GRILLE (slide 13.333" x 7.5")
# --------------------------------------------------------------------------
SW, SH   = Inches(13.333), Inches(7.5)
MARGIN   = Inches(0.85)
CONTENT_W = SW - 2 * MARGIN
TOP_TITLE = Inches(0.72)
TOP_BODY  = Inches(2.05)
FOOT_Y    = Inches(6.86)


# --------------------------------------------------------------------------
# HELPERS BAS NIVEAU
# --------------------------------------------------------------------------
def no_line(shape):
    shape.line.fill.background()
    return shape


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def no_shadow(shape):
    shape.shadow.inherit = False
    return shape


def add_shadow(shape, blur=26, dist=8, alpha=22000, color="0A1A2F"):
    """Ombre portée douce et discrète (XML natif)."""
    spPr = shape._element.spPr
    for tag in ("a:effectLst",):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur*12700)}" dist="{int(dist*12700)}" dir="5400000" '
        'rotWithShape="0">'
        f'<a:srgbClr val="{color}"><a:alpha val="{alpha}"/></a:srgbClr>'
        '</a:outerShdw></a:effectLst>'
    )
    spPr.append(parse_xml(xml))
    return shape


def gradient(shape, c1, c2, angle=0):
    """Dégradé linéaire deux arrêts."""
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    xml = (
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>'
        f'</a:gsLst><a:lin ang="{int(angle*60000)}" scaled="0"/></a:gradFill>'
    )
    # insérer après a:prstGeom / a:custGeom
    geom = spPr.find(qn('a:prstGeom'))
    if geom is None:
        geom = spPr.find(qn('a:custGeom'))
    if geom is not None:
        geom.addnext(parse_xml(xml))
    else:
        spPr.append(parse_xml(xml))
    return shape


def set_corner(shape, value):
    """Rayon d'arrondi d'un ROUNDED_RECTANGLE (0 -> 1)."""
    try:
        shape.adjustments[0] = value
    except Exception:
        pass
    return shape


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE, corner=None):
    s = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    no_line(s)
    no_shadow(s)
    if corner is not None:
        set_corner(s, corner)
    if fill is not None:
        solid(s, fill)
    else:
        s.fill.background()
    s.text_frame.word_wrap = True
    return s


def card(slide, x, y, w, h, fill=WHITE, corner=0.06, shadow=True, border=None):
    s = rect(slide, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, corner)
    if border is not None:
        s.line.color.rgb = border
        s.line.width = Pt(1)
    if shadow:
        add_shadow(s)
    return s


def text(slide, x, y, w, h, runs, size=14, color=GRAPHITE, bold=False,
         font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.25,
         space_after=0, caps=False, spacing=None):
    """runs = str ou liste de (texte, dict d'overrides)."""
    box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        if isinstance(item, tuple):
            content, ov = item
        else:
            content, ov = item, {}
        r = p.add_run()
        r.text = content
        f = r.font
        f.name = ov.get("font", font)
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.color.rgb = ov.get("color", color)
        sp = ov.get("spacing", spacing)
        if sp is not None:
            r.font._rPr.set("spc", str(int(sp * 100)))
        if ov.get("caps", caps):
            r.font._rPr.set("cap", "all")
    return box


def bullets(slide, x, y, w, h, items, size=14, color=SLATE, bullet_color=ACCENT,
            gap=10, bold_lead=False):
    """Liste à puces custom (carrés fins), retourne la textbox."""
    box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(gap)
        lead, rest = (it if isinstance(it, tuple) else (None, it))
        r0 = p.add_run()
        r0.text = "▪  "
        r0.font.size = Pt(size)
        r0.font.color.rgb = bullet_color
        r0.font.name = FONT
        if lead:
            r1 = p.add_run()
            r1.text = lead + " "
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.name = FONT_SEMI
            r1.font.color.rgb = color if not bold_lead else color
        r2 = p.add_run()
        r2.text = rest
        r2.font.size = Pt(size)
        r2.font.name = FONT
        r2.font.color.rgb = color
    return box


def picture_bg(slide, path, prs):
    pic = slide.shapes.add_picture(path, 0, 0, width=prs.slide_width)
    if pic.height < prs.slide_height:
        pic.height = prs.slide_height
        pic.width = int(pic.height * 16 / 9)
    # recadrage centré
    pic.left = int((prs.slide_width - pic.width) / 2)
    pic.top = int((prs.slide_height - pic.height) / 2)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)
    return pic


def scrim(slide, prs, color=INK, alpha=55000):
    """Voile semi-transparent sur toute la slide pour la lisibilité."""
    s = rect(slide, 0, 0, prs.slide_width, prs.slide_height, color)
    fill = s.fill._xPr.find(qn('a:solidFill'))
    clr = fill.find(qn('a:srgbClr'))
    clr.append(parse_xml(
        f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'val="{alpha}"/>'))
    return s
